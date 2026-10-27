# Round-6 docx sync, part B (DD-048): same-machine latency protocol numbers.
# AO 0.424 s median / unfolding 8.257 ms median => ~51x (replaces 1044x,
# which was a cross-machine ratio: AO 1.352 s recorded under 4-worker CPU
# contention; unfolding 1.295 ms from a different deployment machine --
# both retired from ratio use, disclosed).
import docx

SRC = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
d = docx.Document(SRC)

def find(sub, start=0):
    for i, p in enumerate(d.paragraphs[start:], start):
        if sub in p.text:
            return i, p
    return None, None

# ---- abstract: latency phrase -------------------------------------------------
i, p = find('a 1044x speed-up over the committed 1.352 s per-channel AO figure')
if p is not None:
    p.text = p.text.replace(
        'at 1.295 ms median decision latency on a desktop CPU - a 1044x speed-up over '
        'the committed 1.352 s per-channel AO figure',
        'at a same-machine latency ratio of about 51x over the AO solver (AO 0.424 s '
        'per channel, unfolding 8.26 ms per decision, both serialized single-solve wall '
        'clocks in one process on one idle desktop CPU; the previously committed '
        'cross-machine figures - AO 1.352 s recorded under worker contention and a '
        '1.295 ms deployment-machine decision time - are retired from ratio use and '
        'disclosed).')
    print('abstract latency re-footed')
else:
    print('MISS abstract phrase')

# ---- body AO paragraph (P98): median 1.352 sentence ---------------------------
i, p = find('median 1.352 s per channel')
if p is not None:
    p.text = p.text.replace(
        'bps/Hz (mean +/- standard deviation over 500 frozen channels; median 1.352 s per channel)',
        'bps/Hz (mean +/- standard deviation over 500 frozen channels; per-channel solve '
        'time 0.424 s median under the same-machine idle protocol, 0.695 s p95; the earlier '
        '1.352 s figure was recorded inside a four-worker run and carries CPU-contention '
        'inflation - disclosed)')
    print('AO paragraph re-footed')
else:
    print('MISS P98')

# ---- body unfolding paragraph (P118): 1044x rewrite + decomposition -----------
i, p = find('a 1044x speed-up over the 1.352 s per-channel AO reference')
if p is None:
    i, p = find('1044x')
if p is not None:
    p.text = p.text.replace(
        'Deep unfolding is the clear operating point of this regime: 97.41% of the AO '
        'reference (62.53 vs. 64.19 bps/Hz) at 1.295 ms median decision latency on a '
        'desktop CPU, a 1044x speed-up over the 1.352 s per-channel AO reference.',
        'Deep unfolding is the clear operating point of this regime: 97.41% of the AO '
        'reference (62.53 vs. 64.19 bps/Hz) at a same-machine latency ratio of about 51x '
        '(AO 0.424 s vs. unfolding 8.26 ms per decision, serialized single-solve wall '
        'clocks in one process on one idle desktop CPU). The ratio decomposes as about '
        '32x from the iteration budget (median 193 AO iterations versus 6 unrolled '
        'layers) times about 1.6x from cheaper per-operation cost (2.2 ms per AO '
        'iteration vs. 1.4 ms per unrolled layer). The earlier 1044x figure divided an '
        'AO time recorded under 4-worker CPU contention by a decision time measured on a '
        'different machine; it was a cross-machine ratio, is retired from headline use, '
        'and is disclosed (DD-048). A protocol control passes: the untrained forward pass '
        'of the identical graph times at 8.35 ms on the same machine, against 8.26 ms '
        'trained.')
    print('P118 re-footed')
i, p = find('eventual 1044x') or find('1044x speed-up over the 1.352 s per-channel AO reference.')
if p is not None and 'same-machine latency ratio of about 51x' not in p.text:
    p.text = p.text.replace(
        'a 1044x speed-up over the 1.352 s per-channel AO reference.',
        'a same-machine latency ratio of about 51x (AO 0.424 s vs. unfolding 8.26 ms, '
        'serialized single-solve wall clocks on one idle desktop CPU; the retired '
        'cross-machine 1044x came from a contended AO run and a different deployment '
        'machine and is disclosed).')

d.save(SRC)
# verify
a = docx.Document(SRC)
allp = '\n'.join(q.text for q in a.paragraphs)
print('residual 1044x:', allp.count('1044x'), '| 51x mentions:', allp.count('51x'), '| 0.424:', allp.count('0.424'))
