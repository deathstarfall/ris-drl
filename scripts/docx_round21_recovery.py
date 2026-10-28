# Round-21 (DD-062) RECOVERY + auditor fixes.
# The live docx had been silently reverted to a pre-audit version (OneDrive
# sync conflict - old mangled wording resurfaced). Recovery: restore the
# trusted Word-verified snapshot (backups/Report_word_clean_2026-10-28.docx,
# content through DD-061), then apply the auditor's lane-C fixes with strings
# matched to THAT version, then Word-COM verify.
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

LIVE = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
SNAP = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\backups\Report_word_clean_2026-10-28.docx'
FIGS = r'C:\code\ris-drl\paper\figs'

import shutil
shutil.copyfile(SNAP, LIVE)
print('live restored from snapshot')

d = docx.Document(LIVE)
ps = d.paragraphs

def replace_text(p, old, new):
    full = p.text
    assert old in full, 'MISSING: ' + old[:70]
    first = p.runs[0]
    for r in list(p.runs)[1:]:
        r._r.getparent().remove(r._r)
    first.text = full.replace(old, new)

# --- Fix 4 (para 113): quantization numbers from results/bits_sweep.json -------
p = ps[113]
replace_text(
    p,
    'The quantization losses are 2.936 / 0.730 / 0.182 bps/Hz for AO at B = 2/3/4 '
    'against the unquantized 64.193: graceful, monotone degradation in which a '
    '2-bit codebook already retains over 70% of the AO gain over random. Deep '
    "unfolding's own losses are 0.965 / 0.233 / 0.060 bps/Hz against 62.532 - "
    "smaller than AO's at every B.",
    'The quantization losses are 0.976 / 0.243 / 0.061 bps/Hz for AO at B = 2/3/4 '
    'against the unquantized 64.193: graceful, monotone degradation in which a '
    '2-bit codebook already retains over 90% of the AO gain over random (91.1%, '
    'computed against the 11.02 bps/Hz AO-over-random gain). Deep '
    "unfolding's own losses are 0.965 / 0.233 / 0.060 bps/Hz against 62.532 - "
    "smaller than AO's at every B, though only by a hair (0.011/0.010/0.001).")
print('fix4 ok')

# --- Fix 3 wording (para 126): explicit model description ----------------------
p = ps[126]
replace_text(
    p,
    'Deep unfolding is the clear operating point of this regime:',
    'Deep unfolding - an L = 6 unrolled projected-gradient network whose only '
    'trainable parameters are six learned step sizes (6 parameters in code) - '
    'is the clear operating point of this regime:')
print('fix3 wording ok')

# --- Fix 7 (para 129): factored DQN limitation ---------------------------------
p = ps[129]
add7 = (' Sixth, the discrete DQN arms use a factored head that scores each RIS '
        "element's codebook entry independently rather than one joint action over "
        'the codebook; a joint-action DQN over a small codebook is the fair '
        'discrete comparator and remains future work, so the discrete-arm null is '
        'bounded by this design choice.')
assert 'factored head' not in p.text
p.runs[-1].text = p.runs[-1].text + add7
print('fix7 ok')

# --- Fix 8 (para 33): Huang-Mo-Yuen representation sentence --------------------
p = ps[33]
add8 = (' Note that the JSAC 2020 study obtained its positive outcome with a '
        'different input/output representation and reward scale from ours '
        '(state-augmented real/imaginary observations with cosine/sine phase '
        'parameterization versus our raw complex observation with raw-phase or '
        'factored-codebook outputs); representation-matched DRL controls are '
        'therefore flagged as required future work, and our positive-control '
        'probe directly targets this question.')
assert 'representation-matched' not in p.text
p.runs[-1].text = p.runs[-1].text + add8
print('fix8 ok')

# --- Fix 6: figures 2 (training curves) + 3 (AO convergence) before 4.2 --------
from PIL import Image
imgs = [Image.open(rf'{FIGS}\fig_band_{f}.png') for f in
        ('sac', 'td3', 'ddpg', 'dqn_b2', 'dqn_b3', 'dqn_b4')]
w = min(im.width for im in imgs)
imgs = [im.resize((w, int(im.height * w / im.width))) for im in imgs]
canvas = Image.new('RGB', (2 * w, 3 * max(im.height for im in imgs)), 'white')
for k, im in enumerate(imgs):
    canvas.paste(im, ((k % 2) * w, (k // 2) * max(x.height for x in imgs)))
comp = rf'{FIGS}\training_curves_composite.png'
canvas.save(comp)

anchor = ps[108]                      # '4.2' heading
assert anchor.text.startswith('4.2')

def insert_fig_before(anchor_p, img_path, width_in, caption):
    ip = anchor_p.insert_paragraph_before()
    ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ip.add_run().add_picture(img_path, width=Inches(width_in))
    cp = anchor_p.insert_paragraph_before()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.add_run(caption).font.size = Pt(9)

insert_fig_before(anchor, comp, 5.8,
    'Figure 2 Training curves for all six DRL families across the five training '
    'seeds (per-family band plots; solid line is the seed mean, band spans '
    'per-seed min-max). Every family converges to the random-phase floor - the '
    'basis of the null result.')
insert_fig_before(anchor, rf'{FIGS}\fig_ao_convergence.png', 4.8,
    'Figure 3 AO convergence: sum rate versus AO iteration on a representative '
    'test channel (median 193 iterations to tolerance - the origin of the '
    'iteration-budget component of the latency ratio).')

for p in d.paragraphs:
    if p.text.startswith('Figure 2 Achievable mean sum rate'):
        replace_text(p, 'Figure 2', 'Figure 4')
        print('CSI figure renumbered 2 -> 4')
        break

d.save(LIVE)

a = docx.Document(LIVE)
imgs_n = len(a.inline_shapes)
ne = sum(1 for p in a.paragraphs if p.text.strip())
print('VERIFY: paragraphs', len(a.paragraphs), '| non-empty', ne,
      '| tables', len(a.tables), '| images', imgs_n)
for probe in ('0.976 / 0.243 / 0.061', 'six learned step sizes',
              'factored head', 'representation-matched',
              'Figure 2 Training curves', 'Figure 3 AO convergence',
              'Figure 4 Achievable'):
    hit = any(probe in p.text for p in a.paragraphs)
    print(('OK  ' if hit else 'FAIL'), probe)
