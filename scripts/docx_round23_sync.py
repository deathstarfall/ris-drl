# Round-23 (DD-064) v2: fix 3 completion in the docx, with CORRECT anchoring.
# Steps run on a live file that was just restored from the DD-063 snapshot.
# (b1) latency ratio variance wording
# (b2) equal-compute + FLOPs sentences appended to the unfolding paragraph
# (a)  Table 13 inserted physically AFTER the last table (old Table 12)
import docx
from docx.shared import Pt

LIVE = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
d = docx.Document(LIVE)

def find_para(substr):
    for p in d.paragraphs:
        if substr in p.text:
            return p
    raise AssertionError('not found: ' + substr[:50])

def replace_text(p, old, new):
    full = p.text
    assert old in full, 'MISSING: ' + old[:60]
    first = p.runs[0]
    for r in list(p.runs)[1:]:
        r._r.getparent().remove(r._r)
    first.text = full.replace(old, new)

# (b1) ratio variance disclosure
p_op = find_para('is the clear operating point of this regime')
replace_text(p_op,
    'at a same-machine latency ratio of about 51x (AO 0.424 s vs. unfolding '
    '8.26 ms per decision, serialized single-solve wall clocks in one process '
    'on one idle desktop CPU).',
    'at a same-machine latency ratio of roughly 40-50x (frozen protocol run: '
    'AO 0.424 s vs. unfolding 8.26 ms per decision, ratio 51.4x; independent '
    'same-protocol re-measure: AO 0.340 s vs. 8.39 ms, ratio 40.5x; serialized '
    'single-solve wall clocks in one process on one idle desktop CPU).')

# (b2) equal-compute + FLOPs disclosure appended to the decomposition paragraph
p_lat = find_para('The ratio decompos')
add = (" A compute-matched counterpoint table (Table 13) bounds how much of "
       "the gap is architecture versus budget: at six iterations for every "
       'method, truncated AO reaches 61.58 +/- 0.75 and a grid-searched '
       'constant-step projected gradient (step 3.0, selected on a 300-channel '
       'validation hold-out) reaches 62.41 +/- 0.74, while the learned '
       'step sizes reach 62.53 +/- 0.85 - the trained network is the best '
       'six-iteration method, though the margin over a tuned constant step is '
       'only about +0.12 bps/Hz, so the headline ratio is dominated by the '
       'iteration budget, not the architecture. In floating-point terms the '
       'gap is about 48x: one AO decision is roughly 1.37 million complex '
       'multiply-accumulate operations (193 iterations x ~7.1 k each) against '
       'roughly 28.5 thousand for the network (6 layers x ~4.7 k each), a '
       'decomposition of about 32x from iteration count times about 1.5x from '
       'per-iteration cost. Gradient computation is intrinsic to the network: '
       "each layer's update is itself a gradient step, so the per-decision "
       'latency legitimately includes one autograd pass per layer, computed '
       'with the same torch engine the AO reference uses for its gradients.')
if 'Table 13' not in p_lat.text:
    p_lat.runs[-1].text = p_lat.runs[-1].text + add

# (a) build the table + caption at the END of the body, then relocate
tbl = d.add_table(rows=7, cols=3)
tbl.style = 'Table Grid1'
tbl.autofit = False
rows = [
    ('Method', 'Iterations', 'Mean rate (bps/Hz)'),
    ('Random-phase floor (control)', '-', '53.18'),
    ('AO, truncated', '6', '61.58 +/- 0.75'),
    ('Projected gradient (Armijo step)', '6', '59.39 +/- 0.81'),
    ('Projected gradient (constant step 3.0, validation-selected)', '6',
     '62.41 +/- 0.74'),
    ('Learned step sizes (this network)', '6', '62.53 +/- 0.85'),
    ('AO reference (converged)', 'median 193', '64.19 +/- 0.58'),
]
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        run = tbl.rows[ri].cells[ci].paragraphs[0].add_run(val)
        run.font.size = Pt(8)
        run.bold = (ri == 0)
cap = d.add_paragraph()
run = cap.add_run(
    'Table 13 Equal-compute comparison on the frozen test set: every method '
    'limited to six iterations (learned-step-size truncated gradient ascent '
    'versus truncated AO and projected gradient variants; constant-step size '
    'grid-selected on the 300-channel validation hold-out and evaluated once '
    'on test).')
run.font.size = Pt(9)

# relocate to directly after the LAST physical table (old Table 12)
last_tbl = d.tables[11]._tbl     # 12 tables pre-insertion; last = old Table 12
tbl._tbl.getparent().remove(tbl._tbl)
cap._p.getparent().remove(cap._p)
last_tbl.addnext(tbl._tbl)       # table right after old Table 12
tbl._tbl.addprevious(cap._p)     # caption immediately before it

d.save(LIVE)

a = docx.Document(LIVE)
print('tables:', len(a.tables), '| images:', len(a.inline_shapes))
print('tbl12 row2:', [c.text for c in a.tables[11].rows[1].cells])
print('tbl13 row2:', [c.text for c in a.tables[12].rows[1].cells])
caps = [p.text[:8] for p in a.paragraphs if p.text.startswith('Table ')]
print('caption order:', caps)
