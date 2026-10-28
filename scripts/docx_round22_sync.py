# Round-22 (DD-063): second-level auditor lane-C corrections.
# All targets verified against the live docx BEFORE editing (assert in
# replace_text). Jargon policy: DD codes / process jargon must not appear in
# rendered prose (repo comments are fine).
import docx

LIVE = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
d = docx.Document(LIVE)
ps = d.paragraphs

def replace_text(p, old, new):
    full = p.text
    assert old in full, 'MISSING: ' + old[:70]
    first = p.runs[0]
    for r in list(p.runs)[1:]:
        r._r.getparent().remove(r._r)
    first.text = full.replace(old, new)

# 1) para 28: 'repository checker' -> plain manual-verification claim
replace_text(ps[28],
    'All 27 bibliographic entries are verified by the repository checker: DOI resolution',
    'All 27 bibliographic entries were verified as part of the preparation of this '
    'report: DOI resolution')

# 2) para 67: strip 'code convention DD-002'
replace_text(ps[67],
    'applied exactly as implemented (code convention DD-002):',
    'applied exactly as implemented in the released code:')

# 3) para 94: strip DD-041 + artifact path from prose
replace_text(ps[94],
    'selected by validation rate alone (DD-041 protocol; selection artifact '
    'results/unfold/val_protocol_selection.json).',
    'selected by validation rate alone.')

# 4) para 94: strip artifact path in the certified-checkpoint sentence
replace_text(ps[94],
    'are stored in results/baselines/zf_rescore/certified_table_det.json (one '
    'entry per architecture, recording seed 0).',
    'are archived per architecture in the released artifact set.')

# 5) para 94: 'decision log' phrasing
replace_text(ps[94],
    'Hyperparameters, per-run configurations and the full decision log are '
    'committed in configs/ and docs/.',
    'Hyperparameters and per-run configurations are committed in the released '
    'code repository.')

# 6) para 103: neutral wording for the superseded cross-machine pair
replace_text(ps[103],
    'the retired cross-machine pair (1.295 ms / 1.352 s)',
    'an earlier cross-machine measurement pair (1.295 ms / 1.352 s)')

# 7) para 113: v1.1 jargon + stale Figure 2 ref
replace_text(ps[113],
    'as in the report (a robustness sweep over trained DRL checkpoints is left '
    'to the v1.1 campaign). (Table 10; Figure 2)',
    'as in the report (including trained DRL checkpoints in the robustness '
    'sweep is left as future work). (Table 10; Figure 4)')

# 8) para 130: 'retired from headline use'
replace_text(ps[130],
    'it was a cross-machine ratio, is retired from headline use, and is disclosed',
    'it is a cross-machine ratio, is not used in any headline claim, and is '
    'disclosed for completeness')

# 9) quantization paragraph: relative-loss rewording per auditor
p = None
for q in ps:
    if '0.976 / 0.243 / 0.061' in q.text:
        p = q
        break
replace_text(p,
    "smaller than AO's at every B, though only by a hair (0.011/0.010/0.001).",
    'essentially the same relative quantization loss as AO (1.54% vs 1.52% of '
    'the respective unquantized rates at B = 2).')

d.save(LIVE)
print('docx saved')

a = docx.Document(LIVE)
bad = []
for i, p in enumerate(a.paragraphs):
    t = p.text
    for pat in ('DD-0', 'Auditor Round', 'v1.1', 'repository checker', 'retired'):
        if pat.lower() in t.lower():
            bad.append((i, pat, t[:80]))
print('remaining jargon in docx prose:', bad if bad else 'NONE')
print('images:', len(a.inline_shapes), '| tables:', len(a.tables))
