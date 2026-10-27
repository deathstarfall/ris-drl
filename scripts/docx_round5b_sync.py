import docx
from docx.text.paragraph import Paragraph
SRC = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
d = docx.Document(SRC)

def find(sub, start=0):
    for i, p in enumerate(d.paragraphs[start:], start):
        if sub in p.text: return i, p
    return None, None

capP = find('Table 11  Full-observation stress test')[1]
anchor = find('pre-submission stress test closes')[1]
if anchor is None:
    anchor = find('information-side objection')[1]
el = capP._p
nxt = el.getnext()
tbl = nxt if nxt is not None and nxt.tag == W + 'tbl' else None
parent = el.getparent(); parent.remove(el)
if tbl is not None:
    tbl.getparent().remove(tbl)
if tbl is None:  # table was created separately at doc end
    for child in list(d.element.body.iterchildren()):
        if child.tag == W + 'tbl' and 'DQN 4-bit codebook' in child.xml:
            child.getparent().remove(child); tbl = child; break
assert tbl is not None and 'DQN 4-bit codebook' in tbl.xml, 'stress-test table not found'
anchor._p.addnext(tbl)
tbl.addprevious(el)
print('relocated after:', anchor.text[:70])
d.save(SRC)

# verify placement
a = docx.Document(SRC)
body = list(a.element.body.iterchildren())
for k, ch in enumerate(body):
    if ch.tag == W + 'tbl' and 'DQN 4-bit codebook' in ch.xml:
        around = [Paragraph(body[j], a).text[:70] for j in range(max(0, k-2), min(len(body), k+2))
                  if body[j].tag == W + 'p']
        print('table at body pos', k)
        for t in around: print('   ~', t)
