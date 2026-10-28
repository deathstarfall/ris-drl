# Round-17b: reorder tcPr children so w:tcW sits at its schema position
# (tcPr sequence: cnfStyle?, tcW, gridSpan?, ... tcBorders, tcMar, ...). Word's
# strict parser treats out-of-order tcW as corrupt content and repairs on open.
import docx, re
from docx.oxml.ns import qn

DOC = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
d = docx.Document(DOC)

ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd',
         'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark']
moved = 0
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            tcPr = c._tc.find(qn('w:tcPr'))
            if tcPr is None:
                continue
            tcWs = tcPr.findall(qn('w:tcW'))
            if len(tcWs) > 1:            # keep the last (python-docx-managed) one
                for w in tcWs[:-1]:
                    tcPr.remove(w)
                tcWs = [tcWs[-1]]
            if not tcWs:
                continue
            tcW = tcWs[0]
            idx = list(tcPr).index(tcW)
            # first non-tcW child location where tcW must precede
            first_later = None
            for ch in list(tcPr):
                name = ch.tag.split('}')[1]
                if name in ORDER and ORDER.index(name) > ORDER.index('tcW'):
                    first_later = list(tcPr).index(ch)
                    break
            if first_later is not None and idx > first_later:
                tcPr.remove(tcW)
                tcPr.insert(first_later, tcW)
                moved += 1
print('tcW elements moved to schema position:', moved)
d.save(DOC)
print('saved')

a = docx.Document(DOC)
bad = 0
for t in a.tables:
    for row in t.rows:
        for c in row.cells:
            tcPr = c._tc.find(qn('w:tcPr'))
            if tcPr is None: continue
            names = [e.tag.split('}')[1] for e in tcPr]
            if 'tcW' in names and names.index('tcW') > max(
                    [i for i, n in enumerate(names) if n in ('tcBorders', 'tcMar', 'shd')] or [0]):
                bad += 1
print('out-of-order tcPr remaining:', bad, '| tables:', len(a.tables),
      '| images:', len(a.inline_shapes))
