# Round-9 docx sync (DD-051): latency consistency (~1.3 ms -> 8.26 ms same-machine),
# dB-equivalent quantization phrasing dropped (bps/Hz only), CSI random-baseline
# provenance clarified, conclusion narrowed to the tested configuration set.
import docx

SRC = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
d = docx.Document(SRC)

def find(sub, start=0):
    for i, p in enumerate(d.paragraphs[start:], start):
        if sub in p.text:
            return i, p
    return None, None

# 1) Discussion latency consistent with the same-machine protocol
i120, p120 = find('~1.3 ms decision latency')
if p120 is None:
    i120, p120 = find('1.3 ms decision latency')
if p120 is not None:
    p120.text = p120.text.replace('~1.3 ms decision latency', 'an 8.26 ms decision latency (same-machine protocol, Section 5)')
    print('P120 latency fixed')
else:
    print('MISS P120 1.3 ms')

# 2) quantization: drop dB-equivalent, report bps/Hz directly
i130, p130 = find('Table 10  Post-hoc B-bit phase quantization')
if p130 is not None:
    p130.text = ('Table 10  Post-hoc B-bit phase quantization on the frozen test set (500 channels; scope: AO '
                 'reference, deep unfolding and the random control, matching the paper\'s quantization table - '
                 'the DRL arms are evaluated under clean CSI only). Rate losses are reported directly in bps/Hz '
                 'against each method\'s unquantized value (64.193 for AO, 62.532 for unfolding); a rate '
                 'difference is not a dB loss, so no dB conversion is used.')
    print('P130 caption fixed')
i131, p131 = find('0.98 / 0.24 / 0.06 dB-equivalent')
if p131 is None:
    i131, p131 = find('dB-equivalent for AO')
if p131 is not None:
    p131.text = ('The quantization losses are 2.936 / 0.730 / 0.182 bps/Hz for AO at B = 2/3/4 against the '
                 'unquantized 64.193: graceful, monotone degradation in which a 2-bit codebook already retains '
                 'over 70% of the AO gain over random. Deep unfolding\'s own losses are 0.965 / 0.233 / 0.060 '
                 'bps/Hz against 62.532 - smaller than AO\'s at every B. (An earlier "dB-equivalent" description '
                 'of these linear bps/Hz gaps did not use a defined conversion and has been dropped.)')
    print('P131 bps/Hz only')
# any dB-equivalent in table cells
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            if 'dB-equivalent' in c.text:
                c.text = c.text.replace('dB-equivalent', 'bps/Hz (direct)')
                print('cell fixed')

# 3) CSI random baseline provenance (P104 prose mention + caption P105)
ip, pp = None, None
for i, p in enumerate(d.paragraphs):
    if 'channel-state error' in p.text.lower() and '53.172' in p.text:
        ip, pp = i, p
if pp is not None:
    pp.text += (' Provenance note: the 53.172 bps/Hz figure here is this sweep\'s own 200-draw Monte-Carlo '
                'random estimate on the same 500 frozen channels (a fresh sampling realization); the canonical '
                'frozen random baseline used for every delta in this report is 53.178 from the committed '
                'baseline artifact. The two differ by less than 0.01 bps/Hz - independent noise realizations of '
                'the same expectation, not a contradiction - and all headline deltas use the canonical 53.178.')
    print('CSI prose provenance added')
else:
    print('MISS CSI prose (check caption)')
i105, p105 = find('Table 9 Achievable sum rate under channel-state error')
if p105 is not None:
    p105.text = ('Table 9  Achievable sum rate under channel-state error eps (500 frozen test channels; '
                 'decide on h_hat, score on the true channel). The random-phase value here (53.172) is the '
                 'sweep\'s own 200-draw Monte-Carlo realization; the canonical frozen random baseline '
                 '(committed artifact, used for all headline deltas) is 53.178; the two sampling realizations '
                 'of the same expectation differ by < 0.01 bps/Hz.')
    print('CSI caption fixed')

# 4) conclusion narrowed
i127, p127 = find('For RIS phase-shift optimization in a multiuser blocked-direct-link')
if p127 is not None:
    p127.text = p127.text.replace(
        'The null is now unconditional within the regime',
        'The null persists across the tested DRL architectures and observation settings under the specified simulation protocol')
    p127.text = p127.text.replace(
        'unconditionally within the regime',
        'across the tested DRL architectures and observation settings under the specified simulation protocol')
    print('conclusion narrowed check 1')
if p127 is None or 'unconditional' in d.paragraphs[127].text:
    for i, p in enumerate(d.paragraphs):
        if 'unconditional' in p.text:
            k = p.text.find('unconditional')
            print('  residual P'+str(i)+':', p.text[max(0, k - 120):k + 80])

d.save(SRC)
a = docx.Document(SRC)
allp = '\n'.join(q.text for q in a.paragraphs)
print('verify: 1.3 ms left:', allp.count('1.3 ms'), '| dB-equivalent left:', allp.count('dB-equivalent'), '| unconditional left:', allp.count('unconditional'))
