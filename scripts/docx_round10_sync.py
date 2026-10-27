# Round-10 docx sync (DD-052): sixteen-vs-19 clarification, SAC relevance cell,
# Eq (1)/(2) notation and ordering, research-gap claims narrowed to "to our knowledge".
import docx

SRC = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
d = docx.Document(SRC)

def find(sub, start=0):
    for i, p in enumerate(d.paragraphs[start:], start):
        if sub in p.text:
            return i, p
    return None, None

# 1 + 5) Section 2 opener: make the counting and the verification claim exact
ip = None
for i, p in enumerate(d.paragraphs):
    if p.text.startswith('This section reviews sixteen'):
        ip = i
        break
if ip is not None:
    p = d.paragraphs[ip]
    p.text = ('This section reviews sixteen peer-reviewed studies in Tables 1 to 5 that ground the methodology '
              'of this work and motivate the research gap in Section 2.6; the count "sixteen" refers specifically '
              'to Tables 1 to 5. The numbered bibliography additionally contains the three survey papers '
              'summarized in Table 6 (entries [17], [18] and [19]), for 19 entries in total. All 19 bibliographic '
              'entries have been verified against Crossref publisher metadata - DOI resolution plus title match '
              '- for the sixteen corpus studies and the three surveys alike; entries that could not be verified '
              'were dropped entirely rather than annotated.')
    print('Section-2 opener clarified (sixteen vs 19)')

# 2) Table 1 SAC relevance cell
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            if c.text.strip() == 'Foundation of the SAC + codebook arm':
                c.text = 'Foundation of the continuous SAC policy arm (the discrete codebook arms are DQN)'
                print('SAC relevance cell fixed')

# 3) Equations (1) and (2): notation, subscripts and ordering
i60, p60 = find('Θ = diag(e')
if p60 is not None:
    p60.text = ('Θ = diag(e^{jθ₁}, e^{jθ₂}, …, e^{jθ_M}),  θ_m ∈ [0, 2π),  m = 1, …, M     --(1)')
    print('Eq (1) fixed')
i61, p61 = find('Whereby Θ denotes')
if p61 is None:
    i61, p61 = find('Whereby')
if p61 is not None:
    p61.text = ('Θ is the diagonal matrix of RIS phase shifts; θ_m is the phase of the m-th element and M is '
                'the number of surface components. The unit-modulus entries mean each element only reflects, '
                'with no power loss or amplifier gain. The design problem maximizes the achievable sum rate:')
    print('Eq(1) definition paragraph fixed')
i62, p62 = find('maxₓ R(Θ)')
if p62 is None:
    i62, p62 = find('SINRₖ')
if p62 is not None:
    p62.text = ('max over Θ:   R(Θ) = Σ_{k=1..K} log₂(1 + SINR_k(Θ)),   subject to |[Θ]_{m,m}| = 1 for all '
                'm = 1, …, M     --(2)')
    print('Eq (2) fixed')
i63, p63 = None, None
for i, p in enumerate(d.paragraphs):
    if 'SINR_k(Θ) is' not in p.text and p.text.startswith('The control problem is a one-step'):
        i63, p63 = i, p
if p63 is not None:
    p63.text = ('In Eq. (2) the sum runs over the K users and SINR_k(Θ) is user k\'s achievable SINR under the '
                'zero-forcing precoder at the base station with total transmit power P; no per-user power-splitting '
                'heuristic is used. The control problem is a one-step contextual bandit (T = 1): each episode picks '
                'one channel, one phase vector is chosen, the sum rate is observed, and the episode terminates. '
                'Because every episode ends after a single decision, there is no bootstrap term and no discounting.')
    print('Eq(2) definition + bandit note fixed')

# 4) research-gap wording narrowed
gp = find('No one compared them in a fair test')[1]
if gp is not None:
    gp.text = gp.text.replace(
        'No one compared them in a fair test with equal parameters and how much worse results would be for practical ch',
        'To our knowledge, these families have not been compared in one equal-parameter, shared-protocol evaluation covering the practical regimes considered here')
    gp.text = gp.text.replace(
        'There is no paper which is testing how all these approaches behave for various CSI errors and no one tested new user activity.',
        'To our knowledge, prior work has not evaluated these approaches under matched CSI-error sweeps and unseen user activity within a single frozen protocol.')
    print('P20 gaps narrowed')
gp = find('continuous SAC, TD3 and DDPG have never been compared')[1]
if gp is None:
    gp = find('have never been compared')[1]
if gp is not None:
    gp.text = gp.text.replace(
        'continuous SAC, TD3 and DDPG have never been compared with each other and with discrete codebook variants in one frozen',
        'continuous SAC, TD3 and DDPG have not, to our knowledge, been compared with each other or with discrete codebook variants in one frozen')
    gp.text = gp.text.replace(
        'plain baselines - random phases and a no-RIS link - are almost never reported, so published gains cannot be anchored',
        'plain baselines - random phases and a no-RIS link - are rarely reported, so published gains are difficult to anchor')
    print('P46 gaps narrowed')
# also the third gap ("have not been comparatively studied") is fine; check "comparatively"
for p in d.paragraphs:
    if 'have not been comparatively studied' in p.text:
        p.text = p.text.replace('have not been comparatively studied',
                                'have not, to our knowledge, been compared')

d.save(SRC)
a = docx.Document(SRC)
allp = '\n'.join(q.text for q in a.paragraphs)
import re
hits = [m.group(0) for m in re.finditer(r'no one|no paper|never been|almost never', allp, re.I)]
print('residual hard claims:', hits)
print('SAC cell fixed:', 'SAC policy arm' in ' '.join(c.text for t in a.tables for r in t.rows for c in r.cells))
