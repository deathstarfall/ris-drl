# Round-24 phase 2: add training-curve + AO-convergence figures (before 4.2),
# renumber the CSI figure, and insert Table 13 after the last table.
# Runs on the live humanized file AFTER phase 1 (docx_round24_surgical.py).
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

LIVE = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
FIGS = r'C:\code\ris-drl\paper\figs'
d = docx.Document(LIVE)

# locate anchors fresh (indices may shift only via add/remove - none so far)
h42 = None
cap_csi = None
for p in d.paragraphs:
    if p.text.startswith('4.2'):
        h42 = p
    if p.text.startswith('Figure 2 Achievable mean sum rate'):
        cap_csi = p
assert h42 is not None and cap_csi is not None

def insert_fig_before(anchor_p, img_path, width_in, caption):
    ip = anchor_p.insert_paragraph_before()
    ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ip.add_run().add_picture(img_path, width=Inches(width_in))
    cp = anchor_p.insert_paragraph_before()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.add_run(caption).font.size = Pt(9)

insert_fig_before(h42, rf'{FIGS}\training_curves_composite.png', 5.8,
    'Figure 2 Training curves for all six DRL families across the five '
    'training seeds (per-family band plots; solid line is the seed mean, '
    'band spans per-seed min-max). Every family converges to the random-phase '
    'floor - the basis of the null result.')
insert_fig_before(h42, rf'{FIGS}\fig_ao_convergence.png', 4.8,
    'Figure 3 AO convergence: sum rate versus AO iteration on a '
    'representative test channel (median 193 iterations to tolerance - the '
    'origin of the iteration-budget component of the latency ratio).')

# renumber the CSI caption + its in-text ref
first = cap_csi.runs[0]
full = cap_csi.text
for r in list(cap_csi.runs)[1:]:
    r._r.getparent().remove(r._r)
first.text = full.replace('Figure 2 Achievable', 'Figure 4 Achievable', 1)

for p in d.paragraphs:
    if '(Table 10; Figure 2)' in p.text and len(
            p._p.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')) == 0:
        fu = p.text
        r0 = p.runs[0]
        for r in list(p.runs)[1:]:
            r._r.getparent().remove(r._r)
        r0.text = fu.replace('(Table 10; Figure 2)', '(Table 10; Figure 4)', 1)
        print('in-text CSI ref -> Figure 4')
        break

# ---- Table 13 ------------------------------------------------------------------
tbl = d.add_table(rows=7, cols=3)
tbl.style = 'Table Grid1'
tbl.autofit = False
rows = [
    ('Method', 'Iterations', 'Mean rate (bps/Hz)'),
    ('Random-phase floor (control)', '-', '53.18'),
    ('AO, truncated', '6', '61.58 +/- 0.75'),
    ('Projected gradient (Armijo step)', '6', '59.39 +/- 0.81'),
    ('Projected gradient (constant step 3.0, validation-selected)', '6',
     '62.41 +/- 0.74'),
    ('Learned step sizes (this network)', '6', '62.53 +/- 0.85'),
    ('AO reference (converged)', 'median 193', '64.19 +/- 0.58'),
]
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        run = tbl.rows[ri].cells[ci].paragraphs[0].add_run(val)
        run.font.size = Pt(8)
        run.bold = (ri == 0)
cap = d.add_paragraph()
run = cap.add_run(
    'Table 13 Equal-compute comparison on the frozen test set: every method '
    'limited to six iterations (learned-step-size truncated gradient ascent '
    'versus truncated AO and projected gradient variants; constant-step size '
    'grid-selected on the 300-channel validation holdout and evaluated once '
    'on the frozen test set).')
run.font.size = Pt(9)

last_tbl = d.tables[11]._tbl          # 12 tables pre-insertion; last = Table 12
tbl._tbl.getparent().remove(tbl._tbl)
cap._p.getparent().remove(cap._p)
last_tbl.addnext(tbl._tbl)
tbl._tbl.addprevious(cap._p)

d.save(LIVE)

a = docx.Document(LIVE)
ps = a.paragraphs
print('tables:', len(a.tables), '| images:', len(a.inline_shapes))
caps = [p.text[:8] for p in ps if p.text.startswith('Table ')]
print('caption order:', caps)
jargon = [pat for p in ps for pat in ('DD-0', 'DD 0', 'Auditor Round',
          'auditor round', 'v1.1', 'repository checker', 'retired',
          'decision log') if pat.lower() in p.text.lower()]
print('jargon remaining:', jargon if jargon else 'NONE')
for k in ('0.976 / 0.243 / 0.061', 'six learned step sizes', 'factored head',
          'representation-matched', '40-50x', 'about 48x',
          'Table 13 Equal-compute', 'roughly 40-50 times lower'):
    print(('OK  ' if any(k in p.text for p in ps) else 'FAIL'), k)
