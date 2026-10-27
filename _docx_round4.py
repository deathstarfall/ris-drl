# round-4 docx sync (DD-046): reconciled architecture prose, table repairs, scope notes
import docx, shutil

SRC = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report.docx'
BAK = r'C:\Users\user\OneDrive\Desktop\All Assignments\CN project\TeamNo_6_24BDS0166_RIS_6G_Report_pre_round4_sync.docx'
shutil.copyfile(SRC, BAK)
d = docx.Document(SRC)

def sub_in(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new); return True
    if old in p.text:
        p.text = p.text.replace(old, new); return True
    return False

def find(sub, start=0):
    for i, p in enumerate(d.paragraphs[start:], start):
        if sub in p.text: return i, p
    return None, None

# 1) reconcile the six-architecture prose with Table-6 true configs --------------
i, p = find('M1 (DDPG) is a continuous')
if p:
    p.text = ('Six learning architectures are evaluated under one acceptance protocol, in addition to '
        'the unfolding method. M1 (DDPG) is a continuous deterministic actor-critic: fully connected '
        'actor and critic over the 256-value channel observation stack, a 64-dimensional continuous '
        'phase action, an MSE critic loss, Gaussian exploration noise (sigma = 0.1), a 300,000-transition '
        'replay buffer with batch size 256, and Adam updates (actor lr 1e-4, critic lr 1e-3). The '
        'environment is episodic with one decision per episode, so no discounting is applied.')

i, p = find('M2 (SAC + codebook)')
if p:
    p.text = ('M2 (SAC, continuous) acts in the same 64-dimensional continuous phase space: its policy '
        'is a Gaussian network with entropy auto-tuning (initial temperature 0.2; lr 3e-4 for actor, '
        'critic and temperature) - not a Gumbel-Softmax codebook sampler as described in earlier drafts. '
        'The discrete codebook arms are separate DQN variants at 2, 3 and 4-bit resolution (4, 8 and 16 '
        'phase levels per element): the DQN head produces logits over elements x levels (e.g., 512 '
        'logits for 64 elements at 3 bits) and trains with the standard Q-learning objective under '
        'epsilon-greedy exploration. Together these give the six published architectures: SAC, TD3 and '
        'DDPG continuously, plus DQN-2/3/4-bit.')

# 2) bandit paragraph referencing Table 6's discount row --------------------------
i, p = find('gamma reported in Table 6')
if p:
    sub_in(p, 'the gamma reported in Table 6 is inert by construction',
           'no discounting factor is used, because one decision per episode leaves no future term')

# 3) delete the inert discount-factor row from Table 6 ----------------------------
removed = 0
for ti, t in enumerate(d.tables):
    for r in list(t.rows):
        if 'Discount factor' in r.cells[0].text:
            r._element.getparent().remove(r._element); removed += 1
print('discount rows removed:', removed)

# 4) repair Table 5 (survey table): header misaligned, two dead columns -----------
t5 = d.tables[5]
# layout check: data rows have author in col0; header cols were shifted by one
for row in list(t5.rows):
    tcs = row._element.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
    for tc in tcs[-2:]:
        row._element.remove(tc)          # drop the two dead columns
hdr = t5.rows[0]
hdr.cells[0].text = 'Author / Reference'
hdr.cells[1].text = 'Year'
hdr.cells[2].text = 'Venue / Type'
hdr.cells[3].text = 'Coverage'
hdr.cells[4].text = 'Evaluation metrics and key insights'
hdr.cells[5].text = 'Limitations'
print('table5 columns now:', len(t5.rows[0].cells), '| header:', [c.text[:22] for c in hdr.cells])

# 5) quantization scope note -------------------------------------------------------
i, p = find('Table 10')
if p:
    sub_in(p, '(500 channels);',
           '(500 channels; scope: AO reference, deep unfolding and the random control, matching the '
           "paper's quantization table - the DRL arms are evaluated under clean CSI only);")

d.save(SRC)
print('saved')
a = docx.Document(SRC)
t5 = a.tables[5]
for r in t5.rows:
    print(' | '.join(c.text[:40] for c in r.cells))
print('discount rows left:', sum(1 for t in a.tables for r in t.rows if 'Discount' in r.cells[0].text))
print('M2 saclike OK:', 'Gaussian network' in '\n'.join(q.text for q in a.paragraphs))
